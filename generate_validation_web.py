from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Formato de Evaluación de Plataforma Web\nAula Global 360', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Introduction
doc.add_paragraph(
    "Apreciado(a) evaluador(a), el presente instrumento tiene como propósito valorar la herramienta "
    "web Aula Global 360 desde tres aspectos fundamentales: Funcionalidad, Usabilidad y Pertinencia Pedagógica. "
    "Agradecemos su participación diligenciando este formato para la mejora continua del sistema."
)

# Expert Information
doc.add_heading('1. Datos del Evaluador', level=2)
doc.add_paragraph("Nombre completo: ______________________________________________________")
doc.add_paragraph("Cargo / Rol (Ej. Docente, Experto TI): __________________________________")
doc.add_paragraph("Años de experiencia: _______")
doc.add_paragraph("Fecha de evaluación: ___________________")

# Instructions
doc.add_heading('2. Instrucciones', level=2)
doc.add_paragraph(
    "Por favor, marque con una 'X' en la casilla correspondiente su nivel de acuerdo con cada una "
    "de las afirmaciones, utilizando la siguiente escala Likert de 1 a 5:"
)

# Scale Definition
scale_p = doc.add_paragraph()
scale_p.add_run("1: Totalmente en desacuerdo | ").bold = True
scale_p.add_run("2: En desacuerdo | ").bold = True
scale_p.add_run("3: Ni de acuerdo ni en desacuerdo | ").bold = True
scale_p.add_run("4: De acuerdo | ").bold = True
scale_p.add_run("5: Totalmente de acuerdo").bold = True

# Evaluation Criteria (Table)
doc.add_heading('3. Criterios de Evaluación', level=2)

items = [
    ("Criterio 1: Funcionalidad", [
        "1.1. La plataforma procesa los videos de los estudiantes y extrae los fotogramas clave de manera correcta y eficiente.",
        "1.2. El motor de análisis de la inteligencia artificial arroja diagnósticos biomecánicos coherentes con el movimiento observado.",
        "1.3. La plataforma aplica correctamente las configuraciones del profesor (grado, ciclo escolar, tiempo de clase, metodología).",
        "1.4. Las funciones de exportación (descarga de informe individual y unidad didáctica en formato Word) operan correctamente y el documento final está bien estructurado."
    ]),
    ("Criterio 2: Usabilidad", [
        "2.1. La interfaz de usuario es intuitiva y permite navegar fluidamente a través de los tres pasos del sistema (Configuración, Selección, Resultados).",
        "2.2. El diseño estético de la plataforma es visualmente atractivo, moderno y no genera sobrecarga de información.",
        "2.3. La plataforma proporciona notificaciones o indicadores visuales claros sobre su estado actual (ej. tiempos de carga, generación de fotogramas, advertencias).",
        "2.4. Es fácil revertir acciones, regresar a pasos anteriores o modificar la configuración sin perder el hilo del proceso de evaluación."
    ]),
    ("Criterio 3: Pertinencia Pedagógica", [
        "3.1. Las sugerencias y el diagnóstico de la etapa motriz (Inicial, Elemental, Maduro) son coherentes y verdaderamente útiles para la evaluación formativa.",
        "3.2. La unidad didáctica automática se alinea con precisión a las metodologías pedagógicas y materiales físicos que el docente indicó tener disponibles.",
        "3.3. El vocabulario utilizado en los reportes es técnicamente apropiado para el área de la Educación Física y fácilmente comprensible para el profesorado.",
        "3.4. La integración de esta herramienta en la rutina académica representa un avance significativo y real frente a los métodos de planeación tradicionales."
    ])
]

table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Criterio / Descriptor'
hdr_cells[1].text = '1'
hdr_cells[2].text = '2'
hdr_cells[3].text = '3'
hdr_cells[4].text = '4'
hdr_cells[5].text = '5'

# Make header bold
for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            
# Adjust column widths 
for cell in table.columns[0].cells:
    cell.width = Inches(4.5)
for i in range(1, 6):
    for cell in table.columns[i].cells:
        cell.width = Inches(0.4)

for dimension, statements in items:
    # Dimension Header
    row_cells = table.add_row().cells
    row_cells[0].merge(row_cells[5])
    row_cells[0].text = dimension
    for paragraph in row_cells[0].paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            
    # Statements
    for stmt in statements:
        row_cells = table.add_row().cells
        row_cells[0].text = stmt
        for i in range(1, 6):
            row_cells[i].text = ""

# Open Comments
doc.add_heading('4. Observaciones Finales', level=2)
doc.add_paragraph("Describa brevemente las principales fortalezas de la plataforma y mencione cualquier punto crítico o de mejora que haya identificado durante su prueba:\n")
for _ in range(5):
    doc.add_paragraph("_________________________________________________________________________________________________________")

# Signature
doc.add_paragraph("\n\n\n___________________________________\nFirma del Evaluador")

doc.save('d:/Descargas/AulaGlobal360/Formato_Validacion_Plataforma_Web.docx')
print("Formato generado exitosamente.")

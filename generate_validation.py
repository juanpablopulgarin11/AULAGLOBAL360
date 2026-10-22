from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()

# Title
title = doc.add_heading('Formato de Validación por Expertos\nManual de Usuario - Aula Global 360', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Introduction
doc.add_paragraph(
    "Apreciado(a) experto(a), el presente instrumento tiene como propósito evaluar la calidad, "
    "claridad, pertinencia y usabilidad del Manual de Usuario de la plataforma Aula Global 360. "
    "Agradecemos su valiosa colaboración diligenciando este formato."
)

# Expert Information
doc.add_heading('1. Datos del Experto', level=2)
doc.add_paragraph("Nombre completo: ______________________________________________________")
doc.add_paragraph("Profesión / Cargo: ______________________________________________________")
doc.add_paragraph("Años de experiencia relacionada: _______")
doc.add_paragraph("Fecha de evaluación: ___________________")

# Instructions
doc.add_heading('2. Instrucciones', level=2)
doc.add_paragraph(
    "Por favor, marque con una 'X' en la casilla correspondiente su nivel de acuerdo con cada una "
    "de las afirmaciones, utilizando la siguiente escala Likert de 1 a 5:"
)

# Scale Definition
scale_p = doc.add_paragraph()
scale_p.add_run("1: Totalmente en desacuerdo | ").bold = True
scale_p.add_run("2: En desacuerdo | ").bold = True
scale_p.add_run("3: Ni de acuerdo ni en desacuerdo | ").bold = True
scale_p.add_run("4: De acuerdo | ").bold = True
scale_p.add_run("5: Totalmente de acuerdo").bold = True

# Evaluation Criteria (Table)
doc.add_heading('3. Criterios de Evaluación', level=2)

items = [
    ("Dimensión 1: Claridad y Comprensión", [
        "El lenguaje utilizado en el manual es claro, sencillo y adecuado para los docentes.",
        "Las instrucciones paso a paso son fáciles de seguir y comprender.",
        "La explicación sobre los tiempos de espera y la aparición de fotogramas es clara y evita confusiones."
    ]),
    ("Dimensión 2: Diseño Visual y Estructura", [
        "Las capturas de pantalla utilizadas son pertinentes y reflejan fielmente la interfaz del sistema.",
        "La señalización (recuadros rojos y números) en las imágenes facilita ubicar los botones correctamente.",
        "La organización general del documento (títulos, viñetas, secciones) es lógica y ordenada."
    ]),
    ("Dimensión 3: Exactitud Técnica y Completitud", [
        "El manual cubre todas las funciones principales requeridas para operar la plataforma (subida de video, configuración, descargas).",
        "La advertencia sobre no interrumpir el proceso de análisis (clics repetitivos) está bien fundamentada y es necesaria.",
        "Las descripciones de las configuraciones metodológicas y pedagógicas de la IA son precisas."
    ]),
    ("Dimensión 4: Utilidad y Usabilidad", [
        "El manual permite a un usuario sin experiencia previa operar la plataforma sin requerir asistencia técnica adicional.",
        "El documento sirve como una guía de referencia rápida y efectiva ante dudas de funcionamiento.",
        "Considero que el manual cumple con su objetivo de capacitar al docente en el uso de Aula Global 360."
    ])
]

table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Criterio / Afirmación'
hdr_cells[1].text = '1'
hdr_cells[2].text = '2'
hdr_cells[3].text = '3'
hdr_cells[4].text = '4'
hdr_cells[5].text = '5'

# Make header bold
for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            
# Adjust column widths (rough approximation, Word adjusts automatically but this helps)
for cell in table.columns[0].cells:
    cell.width = Inches(4.5)
for i in range(1, 6):
    for cell in table.columns[i].cells:
        cell.width = Inches(0.4)

for dimension, statements in items:
    # Dimension Header
    row_cells = table.add_row().cells
    row_cells[0].merge(row_cells[5])
    row_cells[0].text = dimension
    for paragraph in row_cells[0].paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            
    # Statements
    for stmt in statements:
        row_cells = table.add_row().cells
        row_cells[0].text = stmt
        for i in range(1, 6):
            row_cells[i].text = ""

# Open Comments
doc.add_heading('4. Observaciones Generales y Recomendaciones', level=2)
doc.add_paragraph("¿Qué aspectos considera que deben mejorarse, incluirse o eliminarse del manual?\n")
for _ in range(5):
    doc.add_paragraph("_________________________________________________________________________________________________________")

# Signature
doc.add_paragraph("\n\n\n___________________________________\nFirma del Evaluador / Experto")

doc.save('d:/Descargas/AulaGlobal360/Formato_Validacion_Expertos_Manual.docx')
print("Formato generado exitosamente.")

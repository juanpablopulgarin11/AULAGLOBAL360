"""Documentos Word (.docx) reales con python-docx.

Reemplazan a ``exportDiagnosticoToWord`` y ``exportToWord`` del JS, que descargaban HTML con
extensión ``.doc``. Se conserva la estructura y el formato institucional (carta vertical,
márgenes de 1.8 cm, Calibri 10 pt, encabezados azul #0284C7).

Las funciones reciben ``dict`` (el contrato ``Diagnostico`` y la unidad de
``biomecanica.didactica``) para poder probarse sin base de datos.
"""
from __future__ import annotations

import datetime as dt
import io
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = "0284C7"
AZUL_OSCURO = "0369A1"
GRIS_FONDO = "F1F5F9"
GRIS_SUAVE = "F8FAFC"
VERDE, ROJO = "059669", "DC2626"


# --------------------------------------------------------------------------------------
# Utilidades de formato
# --------------------------------------------------------------------------------------
def _documento() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)          # carta
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(1.8))
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    estilo.paragraph_format.space_after = Pt(2)
    return doc


def _sombrear(celda, color: str) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _escribir(celda, texto: str = "", negrita: bool = False, color: Optional[str] = None, tam: float = 9.5,
              cursiva: bool = False, centrado: bool = False, etiqueta: Optional[str] = None) -> None:
    """Reemplaza el contenido de la celda. ``etiqueta`` se escribe en negrita antes del texto."""
    p = celda.paragraphs[0]
    p.text = ""
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if etiqueta:
        r = p.add_run(etiqueta + " ")
        r.bold, r.font.size = True, Pt(tam)
    r = p.add_run(texto or "")
    r.bold, r.italic, r.font.size = negrita, cursiva, Pt(tam)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def _encabezado(celda, texto: str, fondo: str = AZUL, centrado: bool = False) -> None:
    _sombrear(celda, fondo)
    _escribir(celda, texto.upper(), negrita=True, color="FFFFFF" if fondo == AZUL else "334155", tam=10, centrado=centrado)


def _subencabezado(celda, texto: str, centrado: bool = False) -> None:
    _sombrear(celda, GRIS_FONDO)
    _escribir(celda, texto, negrita=True, color="334155", tam=9, centrado=centrado)


def _tabla(doc: Document, filas: int, columnas: int, anchos: Optional[Sequence[float]] = None):
    t = doc.add_table(rows=filas, cols=columnas)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    ancho = OxmlElement("w:tblW")          # 100 % del ancho útil de la página
    ancho.set(qn("w:type"), "pct")
    ancho.set(qn("w:w"), "5000")
    for previo in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(previo)
    tbl_pr.append(ancho)
    if anchos:
        for fila in t.rows:
            for celda, ancho in zip(fila.cells, anchos):
                celda.width = Cm(ancho)
    return t


def _fila_unida(tabla, fila: int):
    celdas = tabla.rows[fila].cells
    return celdas[0].merge(celdas[-1]) if len(celdas) > 1 else celdas[0]


def _lista(celda, items: Iterable[str], cursiva: bool = False, comillas: bool = False) -> None:
    primero = True
    for item in items:
        p = celda.paragraphs[0] if primero else celda.add_paragraph()
        primero = False
        p.text = ""
        r = p.add_run(f"• {'“' + item + '”' if comillas else item}")
        r.italic, r.font.size = cursiva, Pt(9.5)


def _parrafo(celda, partes: Sequence[tuple], tam: float = 9.5) -> None:
    """Párrafo nuevo en la celda con trozos ``(texto, negrita, cursiva)`` del mismo tamaño."""
    p = celda.add_paragraph()
    for texto, negrita, cursiva in partes:
        r = p.add_run(texto)
        r.bold, r.italic, r.font.size = negrita, cursiva, Pt(tam)


def _espacio(doc: Document, pt: float = 6) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(pt)


def _firmas(doc: Document, izquierda: Sequence[str], derecha: Sequence[str]) -> None:
    _espacio(doc, 24)
    t = doc.add_table(rows=1, cols=2)
    for celda, lineas in zip(t.rows[0].cells, (izquierda, derecha)):
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("_" * 40 + "\n")
        r = p.add_run(lineas[0] + "\n")
        r.bold = True
        for extra in lineas[1:]:
            r = p.add_run(extra)
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string("64748B")


def _bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _limpiar(texto: str) -> str:
    """Quita el marcado **negrita** del resumen (en Word se muestra texto plano)."""
    return (texto or "").replace("**", "")


# --------------------------------------------------------------------------------------
# Reporte del estudiante
# --------------------------------------------------------------------------------------
def reporte_estudiante(d: Mapping[str, Any], estudiante: str = "", grupo: str = "", docente: str = "",
                       institucion: str = "", fecha: Optional[dt.date] = None,
                       imagenes: Sequence[bytes] = (), advertencias: Sequence[str] = ()) -> bytes:
    """Informe biomecánico para el acudiente y el historial (equivalente a ``exportDiagnosticoToWord``)."""
    fecha = fecha or dt.date.today()
    doc = _documento()

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = titulo.add_run("INFORME DE EVALUACIÓN BIOMECÁNICA HMB")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(14), RGBColor.from_string(AZUL_OSCURO)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"{institucion or 'AULA GLOBAL 360'} · Batería Validada (Dialnet 7925607) · Fecha: {fecha:%d/%m/%Y}")
    r.font.size, r.font.color.rgb = Pt(9), RGBColor.from_string("64748B")

    # 1. Datos generales
    t = _tabla(doc, 6 if estudiante else 5, 2, (8.9, 8.9))
    _encabezado(_fila_unida(t, 0), "1. Datos generales y resultado evolutivo")
    fila = 1
    if estudiante:
        _escribir(t.cell(fila, 0), estudiante, etiqueta="Estudiante:")
        _escribir(t.cell(fila, 1), grupo or "—", etiqueta="Grupo:")
        fila += 1
    _escribir(t.cell(fila, 0), d["habilidad_detectada"].upper(), etiqueta="Habilidad evaluada:")
    _escribir(t.cell(fila, 1), d.get("componente_hmb") or "", etiqueta="Componente:")
    _escribir(t.cell(fila + 1, 0), f"{d.get('puntaje_obtenido')} ({d.get('porcentaje_madurez')}% de madurez)", etiqueta="Puntuación Batería HMB:")
    _escribir(t.cell(fila + 1, 1), (d.get("estadio_gallahue") or "").upper(), negrita=True, color=AZUL, etiqueta="Estadio (Gallahue):")
    _escribir(t.cell(fila + 2, 0), d.get("edad_calibrada") or "", etiqueta="Grado / edad:")
    _escribir(t.cell(fila + 2, 1), "Detección automática por cinemática" if d.get("es_deteccion_automatica") else "Evaluación dirigida",
              etiqueta="Modalidad:")
    marco = _fila_unida(t, fila + 3)
    _sombrear(marco, GRIS_SUAVE)
    _escribir(marco, "Batería de Habilidades Motrices Básicas para Niños entre 5 y 11 Años (González Palacio, Montoya Grisales, "
                     "Cardona, Marín & Muñoz, 2021 · Dialnet 7925607) y Estadios de Desarrollo Motor (David L. Gallahue).",
              tam=8.5, color="475569", etiqueta="Marco científico:")
    _espacio(doc)

    # 2. Telemetría
    tel = d.get("telemetria_medida") or {}
    filas_tel = [
        ("Flexión mínima de rodilla (recobro / carga)", f"{tel.get('minKneeAngle', '—')}°"),
        ("Ángulo medio de codos (braceo)", f"{tel.get('avgElbowAngle', '—')}°"),
        ("Inclinación del tronco respecto a la vertical", f"{tel.get('avgTrunkAngle', '—')}°"),
        ("Apertura angular de cadera / zancada", f"{tel.get('maxHipAngle', '—')}°"),
        ("Fase aérea / vuelo", "DETECTADA" if tel.get("flightDetected") else "NO EVIDENTE / APOYO EN SUELO"),
        ("Simetría bilateral", f"{tel.get('symmetryScore', '—')}%"),
    ]
    t = _tabla(doc, len(filas_tel) + 2, 2, (11, 6.8))
    _encabezado(_fila_unida(t, 0), "2. Telemetría cinemática articular (MediaPipe Pose · 33 puntos)")
    _subencabezado(t.cell(1, 0), "Variable articular medida")
    _subencabezado(t.cell(1, 1), "Medición obtenida", centrado=True)
    for i, (variable, valor) in enumerate(filas_tel, start=2):
        _escribir(t.cell(i, 0), variable, negrita=True)
        _escribir(t.cell(i, 1), valor, centrado=True)
    _espacio(doc)

    # 3. Criterios
    criterios = d.get("criterios") or []
    t = _tabla(doc, len(criterios) + 2, 2, (13.8, 4))
    _encabezado(_fila_unida(t, 0), "3. Criterios biomecánicos contrastados (0 / 1)")
    _subencabezado(t.cell(1, 0), "Criterio técnico evaluado y medición")
    _subencabezado(t.cell(1, 1), "Resultado", centrado=True)
    for i, c in enumerate(criterios, start=2):
        celda = t.cell(i, 0)
        _escribir(celda, f"{i - 1}. {c['criterio']}", negrita=True)
        detalles = [f"Fase: {c.get('fase') or 'Ejecución'}"]
        if c.get("medido"):
            detalles.append(f"Medido: {c['medido']} | Umbral: {c.get('umbral') or '—'}")
        if c.get("observacion"):
            detalles.append(f"Observación: {c['observacion']}")
        for linea in detalles:
            r = celda.add_paragraph().add_run(linea)
            r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string("475569")
        logrado = c.get("puntaje") == 1
        _sombrear(t.cell(i, 1), "ECFDF5" if logrado else "FEF2F2")
        _escribir(t.cell(i, 1), "✓ LOGRADO" if logrado else "✗ EN PROCESO", negrita=True, color=VERDE if logrado else ROJO, centrado=True)
    _espacio(doc)

    # 4. Síntesis
    t = _tabla(doc, 2, 1, (17.8,))
    _encabezado(t.cell(0, 0), "4. Síntesis biomecánica y anomalías observadas")
    celda = t.cell(1, 0)
    _escribir(celda, _limpiar(d.get("resumen_biomecanico", "")))
    p = celda.add_paragraph()
    r = p.add_run("ANOMALÍAS DETECTADAS EN LA CADENA CINÉTICA:")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor.from_string("B91C1C")
    for e in d.get("errores_criticos") or []:
        _parrafo(celda, [(f"• {e['error']}: ", True, False), (e.get("impacto_biomecanico", ""), False, False)], tam=9)
    for aviso in advertencias:
        r = celda.add_paragraph().add_run(f"Nota sobre la grabación: {aviso}")
        r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor.from_string("B45309")
    _espacio(doc)

    # 5. Consignas
    t = _tabla(doc, 2, 1, (17.8,))
    _encabezado(t.cell(0, 0), "5. Consignas verbales para el estudiante (\"El lenguaje del profe\")")
    _lista(t.cell(1, 0), d.get("frases_profe") or [], cursiva=True, comillas=True)

    # 6. Evidencia (fotogramas con esqueleto)
    if imagenes:
        _espacio(doc)
        cols = 4
        filas = (len(imagenes) + cols - 1) // cols
        t = _tabla(doc, filas + 1, cols)
        _encabezado(_fila_unida(t, 0), "6. Evidencia: fotogramas analizados")
        for k, img in enumerate(imagenes):
            celda = t.cell(1 + k // cols, k % cols)
            celda.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            celda.paragraphs[0].add_run().add_picture(io.BytesIO(img), width=Cm(4.1))

    _firmas(doc, ["Firma docente evaluador", docente or "Docente de Educación Física"],
            ["Firma acudiente / padre de familia", "C.C. ________________________"])
    return _bytes(doc)


# --------------------------------------------------------------------------------------
# Unidad didáctica
# --------------------------------------------------------------------------------------
def unidad_didactica(u: Mapping[str, Any], sesiones: Optional[Sequence[Mapping[str, Any]]] = None,
                     institucion: str = "", docente: str = "") -> bytes:
    """Planeación del período con todas sus sesiones (equivalente a ``exportToWord``).

    ``sesiones`` permite pasar las sesiones editadas por el docente; si no, se usan las de ``u``.
    """
    sesiones = list(sesiones if sesiones is not None else u["clases_secuencia"])
    doc = _documento()

    # 1. Encabezado institucional
    t = _tabla(doc, 7, 4, (4.9, 4.6, 4.6, 3.7))
    _encabezado(_fila_unida(t, 0), "Unidad didáctica · Formato institucional de planeación curricular", centrado=True)
    inst = _fila_unida(t, 1)
    _subencabezado(inst, institucion or u.get("institucion", ""), centrado=True)
    _escribir(t.cell(2, 0), u["area"], etiqueta="ÁREA:")
    _escribir(t.cell(2, 1), u["ciclo"], etiqueta="CICLO:")
    _escribir(t.cell(2, 2), u["grado"], etiqueta="GRADO:")
    _escribir(t.cell(2, 3), u["periodo"], etiqueta="PERÍODO:")
    _escribir(t.cell(3, 0), docente or u.get("docente", ""), etiqueta="DOCENTE:")
    _escribir(t.cell(3, 1), u["anio"], etiqueta="AÑO:")
    _escribir(t.cell(3, 2), u["jornada"], etiqueta="JORNADA:")
    _escribir(t.cell(3, 3), u["duracion_clase"], etiqueta="DURACIÓN:")
    _escribir(_fila_unida(t, 4), f"{u['tema']} · (Secuencia progresiva de {len(sesiones)} clases planificadas)", etiqueta="UNIDAD TEMÁTICA:")
    _escribir(_fila_unida(t, 5), u["lugar"], etiqueta="LUGAR / INSTALACIÓN:")
    _escribir(_fila_unida(t, 6), f"{u['materiales']} · Formato: {u['formato']} · Metodología: {u['metodologia']}", etiqueta="MATERIALES:")
    _espacio(doc)

    # 2. Pregunta, objetivos, lineamientos e indicadores
    t = _tabla(doc, 9, 3, (5.9, 6, 5.9))
    _encabezado(_fila_unida(t, 0), "Pregunta problematizadora y objetivos de aprendizaje")
    _escribir(_fila_unida(t, 1), u["pregunta_problematizadora"], cursiva=True, etiqueta="Pregunta problematizadora:")
    _escribir(_fila_unida(t, 2), u["objetivo_general"], etiqueta="Objetivo general:")
    obj = _fila_unida(t, 3)
    _lista(obj, u["objetivos_especificos"])
    _encabezado(_fila_unida(t, 4), "Lineamientos curriculares / orientaciones pedagógicas (MEN Colombia)")
    for j, (titulo, clave) in enumerate((("Competencia motriz", "motriz"), ("Competencia expresivo-corporal", "expresivo_corporal"),
                                          ("Competencia axiológica-corporal", "axilogica_corporal"))):
        _subencabezado(t.cell(5, j), titulo)
        _escribir(t.cell(6, j), u["estandares"][clave])
    for j, (titulo, clave) in enumerate((("SABER | Cognitivo", "saber"), ("HACER | Procedimental", "hacer"), ("SER | Actitudinal", "ser"))):
        _subencabezado(t.cell(7, j), titulo)
        _escribir(t.cell(8, j), u["indicadores"][clave])
    _espacio(doc)

    # 3. Secuencia de sesiones
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = titulo.add_run(f"SECUENCIA DIDÁCTICA Y MATRIZ DE PROGRESIÓN ({len(sesiones)} SESIONES)")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), RGBColor.from_string(AZUL_OSCURO)
    for s in sesiones:
        t = _tabla(doc, 6, 2, (4.9, 12.9))
        cab = _fila_unida(t, 0)
        _encabezado(cab, f"Sesión {s['numero']} de {len(sesiones)}: {s['titulo']}")
        r = cab.paragraphs[0].add_run(f"   {s.get('fase_pedagogica', '')} · {u['duracion_clase']}")
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string("FFFFFF")
        _subencabezado(t.cell(1, 0), "OBJETIVO")
        _escribir(t.cell(1, 1), s["objetivo"])
        _subencabezado(t.cell(2, 0), f"PARTE INICIAL ({s.get('minutos_inicial', '')} min)\nActivación y movilidad")
        celda = t.cell(2, 1)
        if s.get("es_refuerzo") and s.get("error_reforzado"):
            _sombrear(celda, "EFF6FF")
            _escribir(celda, f"En la evaluación se identificó «{s['error_reforzado']}». Hoy se trabaja con énfasis "
                             "prioritario para corregir esta falencia.", color="1E40AF", etiqueta="🎯 Refuerzo dirigido:")
            _parrafo(celda, [(s["actividad_inicial"], False, False)])
        else:
            _escribir(celda, s["actividad_inicial"])
        _subencabezado(t.cell(3, 0), f"PARTE CENTRAL ({s.get('minutos_central', '')} min)\nDesarrollo y tareas motrices")
        celda = t.cell(3, 1)
        _escribir(celda, s.get("distribucion", ""), etiqueta=f"1. Montaje ({u['formato']}):")
        _parrafo(celda, [("2. Desarrollo: ", True, False), (s["actividad_central"], False, False)])
        _parrafo(celda, [("3. Consigna clave: ", True, False), (f"“{s.get('consigna', '')}”", False, True)])
        _subencabezado(t.cell(4, 0), f"PARTE FINAL ({s.get('minutos_final', '')} min)\nVuelta a la calma")
        _escribir(t.cell(4, 1), s["actividad_final"])
        _subencabezado(t.cell(5, 0), "INDICADOR DE EVALUACIÓN")
        _escribir(t.cell(5, 1), s.get("criterio_eval", ""), etiqueta="Criterio de logro:")
        _espacio(doc, 4)

    # 4. Complementos
    filas = [
        ("TAREA Y REPASO EXTRACURRICULAR", u["tarea_extracurricular"]),
        ("MÉTODOS DE ENSEÑANZA", u["metodos_ensenanza"]),
        ("ESTILO DE ENSEÑANZA", u["estilo_ensenanza"]),
        ("ADAPTACIONES RAZONABLES (PIAR / DUA)", u["adaptaciones_piar"]),
        ("EVALUACIÓN FORMATIVA", u["evaluacion"]),
        ("REFLEXIÓN PEDAGÓGICA", u["reflexion_pedagogica"]),
        ("CONSIGNAS CLAVE (\"El lenguaje del profe\")", None),
        ("BIBLIOGRAFÍA Y REFERENTES CURRICULARES", u["bibliografia"]),
    ]
    t = _tabla(doc, len(filas) + 1, 2, (5.7, 12.1))
    _encabezado(_fila_unida(t, 0), "Lineamientos metodológicos, inclusión DUA/PIAR y sistema evaluativo")
    for i, (titulo, texto) in enumerate(filas, start=1):
        _subencabezado(t.cell(i, 0), titulo)
        if texto is None:
            _lista(t.cell(i, 1), u["retroalimentacion_tips"], cursiva=True, comillas=True)
        else:
            _escribir(t.cell(i, 1), texto)

    _firmas(doc, ["Firma del docente titular de Educación Física", "C.C. ________________________"],
            ["Firma de coordinación académica / directiva", "Institución educativa"])
    return _bytes(doc)


def nombre_archivo(*partes: str, extension: str = "docx") -> str:
    """``Reporte_Juan_Perez_Salto_Horizontal.docx`` (sin tildes ni espacios)."""
    import unicodedata

    base = "_".join(p for p in partes if p)
    base = unicodedata.normalize("NFKD", base).encode("ascii", "ignore").decode()
    base = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in base)
    while "__" in base:
        base = base.replace("__", "_")
    return f"{base.strip('_') or 'documento'}.{extension}"

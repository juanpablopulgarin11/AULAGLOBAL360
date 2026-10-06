import io

import pytest
from docx import Document

from apps.catalogo.models import Habilidad
from apps.cuentas.models import Institucion
from apps.estudiantes.models import Estudiante, Grupo
from apps.evaluaciones.models import Evaluacion, EvaluacionGrupal
from apps.evaluaciones.servicios import diagnosticar_y_guardar
from apps.planeacion.servicios import preferencias, unidad_para_evaluacion, unidad_para_grupo
from apps.reportes.servicios import documento_de_unidad, reporte_de_evaluacion
from apps.reportes.word import nombre_archivo


def _pose():
    c = {11: (0.55, 0.30), 12: (0.45, 0.30), 13: (0.57, 0.42), 14: (0.43, 0.42), 15: (0.58, 0.53), 16: (0.42, 0.53),
         23: (0.53, 0.55), 24: (0.47, 0.55), 25: (0.53, 0.72), 26: (0.47, 0.72), 27: (0.53, 0.90), 28: (0.47, 0.90)}
    return [{"x": c.get(i, (0.5, 0.15))[0], "y": c.get(i, (0.5, 0.15))[1], "z": 0.0} for i in range(33)]


def _texto(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                partes.append(celda.text)
    return "\n".join(partes)


@pytest.fixture
def evaluacion(catalogo, docente):
    inst = Institucion.objects.create(nombre="IE La Esperanza")
    docente.institucion = inst
    docente.save()
    grupo = Grupo.objects.create(institucion=inst, docente=docente, nombre="3ºA", grado="8_anos", anio=2026)
    est = Estudiante.objects.create(grupo=grupo, nombres="María José", apellidos="Gómez")
    ev = Evaluacion.objects.create(docente=docente, estudiante=est, grado="8_anos",
                                   habilidad_solicitada=Habilidad.objects.get(codigo="carrera"))
    diagnosticar_y_guardar(ev, [{"timestampNum": k * 0.2, "landmarks": _pose()} for k in range(8)])
    return ev


@pytest.mark.django_db
def test_unidad_individual_prioriza_falencias(evaluacion, docente):
    prefs = preferencias(materiales=["Conos", "Balones"], periodo=2, duracion_min=45, total_clases=8)
    ud = unidad_para_evaluacion(evaluacion, docente, prefs)
    sesiones = list(ud.sesiones.all())
    assert ud.total_clases == 8 and len(sesiones) == 8 and ud.materiales == ["Conos", "Balones"]
    assert sesiones[0].plantilla.orden == 1 and not sesiones[0].es_refuerzo
    assert any(s.es_refuerzo and s.error_reforzado for s in sesiones[1:3])   # una postura estática falla criterios de carrera
    assert ud.contenido["duraciones"]["inicial"] == "9 minutos"
    assert "Conos, Balones" in sesiones[0].distribucion or "Conos, Balones" in ud.contenido["materiales"]


@pytest.mark.django_db
def test_unidad_rechaza_evaluacion_sin_diagnostico(catalogo, docente):
    ev = Evaluacion.objects.create(docente=docente)
    with pytest.raises(ValueError):
        unidad_para_evaluacion(ev, docente, preferencias())


@pytest.mark.django_db
def test_unidad_grupal_usa_la_habilidad_mas_debil(evaluacion, docente):
    grupal = EvaluacionGrupal.objects.create(grupo=evaluacion.estudiante.grupo, docente=docente, estudiantes_objetivo=2)
    evaluacion.evaluacion_grupal = grupal
    evaluacion.save()
    ud = unidad_para_grupo(grupal, docente, preferencias(total_clases=4))
    assert ud.habilidad.codigo == "carrera" and ud.evaluacion_grupal == grupal
    assert ud.contenido["grado"] == "Salón Completo (Heterogéneo)"
    assert ud.sesiones.filter(es_refuerzo=True).exists()


@pytest.mark.django_db
def test_reporte_word_del_estudiante(evaluacion):
    contenido, nombre = reporte_de_evaluacion(evaluacion)
    texto = _texto(contenido)
    assert nombre == "Reporte_Maria_Jose_Gomez_Carrera.docx"
    assert "INFORME DE EVALUACIÓN BIOMECÁNICA HMB" in texto and "María José Gómez" in texto
    assert "IE La Esperanza" in texto and "EN PROCESO" in texto
    assert "WASM" not in texto                                   # textos del motor en servidor


@pytest.mark.django_db
def test_word_de_la_unidad_refleja_ediciones(evaluacion, docente):
    ud = unidad_para_evaluacion(evaluacion, docente, preferencias(total_clases=3))
    s = ud.sesiones.get(numero=2)
    s.objetivo = "Objetivo ajustado por la docente"
    s.save()
    contenido, nombre = documento_de_unidad(ud)
    texto = _texto(contenido)
    assert nombre.startswith("Unidad_Didactica_Periodo1_3Clases_Carrera")
    assert "Objetivo ajustado por la docente" in texto and "SESIÓN 3 DE 3" in texto.upper()


def test_nombre_de_archivo_seguro():
    assert nombre_archivo("Reporte", "Ñandú José/../x", "Salto Horizontal") == "Reporte_Nandu_Jose_x_Salto_Horizontal.docx"

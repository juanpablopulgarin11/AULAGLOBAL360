"""Flujo completo de la interfaz web: registro → salón → evaluar video → resultado → planeación → Word."""
from __future__ import annotations

import cv2
import numpy as np
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.estudiantes.models import Estudiante, Grupo
from apps.evaluaciones import tasks
from apps.evaluaciones.models import Evaluacion, EvaluacionGrupal
from apps.planeacion.models import UnidadDidactica

from tests.test_extraccion import detector_falso, escribir_video

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def cliente(client, docente):
    client.force_login(docente)
    return client


@pytest.fixture(autouse=True)
def sin_mediapipe(monkeypatch):
    monkeypatch.setattr(tasks, "obtener_detector", lambda: detector_falso)


def video_subida(tmp_path, nombre="salto.mp4"):
    return SimpleUploadedFile(nombre, escribir_video(tmp_path / nombre).read_bytes(), content_type="video/mp4")


def evaluar(cliente, tmp_path, django_capture_on_commit_callbacks, **datos):
    datos = {"habilidad": "salto", "grado": "7_anos", "motor": "local", **datos}
    with django_capture_on_commit_callbacks(execute=True):
        r = cliente.post(reverse("evaluar"), {**datos, "archivo": video_subida(tmp_path)})
    assert r.status_code == 302, r.content.decode()[:2000]
    return Evaluacion.objects.latest("pk")


@pytest.mark.django_db
def test_paginas_publicas_y_login_requerido(client):
    assert client.get(reverse("inicio")).status_code == 200
    assert "9 Habilidades" in client.get(reverse("inicio")).content.decode()
    assert client.get(reverse("ingresar")).status_code == 200
    for nombre in ("panel", "evaluar", "evaluaciones", "grupos", "unidades"):
        r = client.get(reverse(nombre))
        assert r.status_code == 302 and "/cuentas/ingresar/" in r["Location"], nombre


@pytest.mark.django_db
def test_registro_crea_docente_con_institucion(client):
    r = client.post(reverse("registro"), {"username": "lucia", "first_name": "Lucía", "last_name": "Ríos",
                                          "email": "lucia@colegio.edu.co", "institucion": "IE San José",
                                          "password1": "Clave-segura-2026", "password2": "Clave-segura-2026"})
    assert r.status_code == 302 and r["Location"] == reverse("panel")
    assert client.get(reverse("panel")).status_code == 200
    from apps.cuentas.models import Docente
    assert Docente.objects.get(username="lucia").institucion.nombre == "IE San José"


@pytest.mark.django_db
def test_flujo_completo_individual(cliente, catalogo, docente, tmp_path, django_capture_on_commit_callbacks):
    assert cliente.get(reverse("evaluar")).status_code == 200

    ev = evaluar(cliente, tmp_path, django_capture_on_commit_callbacks)
    assert ev.estado == "lista" and ev.habilidad_detectada.codigo == "salto"

    detalle = cliente.get(reverse("evaluacion", args=[ev.pk]))
    html = detalle.content.decode()
    assert detalle.status_code == 200 and "Estadio" in html and "Ver criterios evaluados" in html
    assert cliente.get(reverse("evaluacion_estado", args=[ev.pk])).json()["estado"] == "lista"

    img = cliente.get(reverse("fotograma", args=[ev.pk, 1, "esqueleto"]))
    assert img.status_code == 200 and img["Content-Type"] == "image/jpeg"

    rep = cliente.get(reverse("reporte", args=[ev.pk]))
    assert rep.status_code == 200 and rep["Content-Type"] == DOCX and rep.content[:2] == b"PK"

    r = cliente.post(reverse("unidad_crear", args=[ev.pk]), {"formato": "Cuento Motor", "metodologia": "Descubrimiento Guiado",
                                                              "periodo": 3, "duracion_min": 55, "total_clases": 4,
                                                              "materiales": ["Conos", "Colchonetas"]})
    ud = UnidadDidactica.objects.get()
    assert r.status_code == 302 and r["Location"] == reverse("unidad", args=[ud.pk])
    assert ud.total_clases == 4 and ud.periodo == 3 and ud.materiales == ["Conos", "Colchonetas"]
    assert cliente.get(reverse("unidad", args=[ud.pk])).status_code == 200

    r = cliente.post(reverse("sesion_editar", args=[ud.pk, 2]), {
        "titulo": "Sesión ajustada", "objetivo": "Nuevo objetivo", "distribucion": "d", "actividad_inicial": "a",
        "actividad_central": "c", "actividad_final": "f", "consigna": "¡Vamos!", "criterio_eval": "e"})
    assert r.status_code == 302 and ud.sesiones.get(numero=2).titulo == "Sesión ajustada"

    doc = cliente.get(reverse("unidad_docx", args=[ud.pk]))
    assert doc.status_code == 200 and doc["Content-Type"] == DOCX
    assert cliente.get(reverse("unidades")).status_code == 200
    assert cliente.get(reverse("evaluaciones") + "?estadio=Inicial").status_code == 200
    assert cliente.get(reverse("panel")).status_code == 200
    comprobar_word_con_fotos_3d(cliente, ud)


def comprobar_word_con_fotos_3d(cliente, ud):
    """El navegador toma las fotos del montaje 3D (capturas3d.js) y el Word las pone bajo cada sesión."""
    import io

    from docx import Document

    escenas = cliente.get(reverse("unidad_escenas", args=[ud.pk])).json()["sesiones"]
    assert [s["numero"] for s in escenas] == [1, 2, 3, 4]
    assert all(s["escena"]["pasos"] and s["escena"]["elementos"] for s in escenas)

    ok, jpg = cv2.imencode(".jpg", np.full((75, 120, 3), 200, np.uint8))
    foto = lambda nombre: SimpleUploadedFile(nombre, jpg.tobytes(), content_type="image/jpeg")  # noqa: E731
    s1 = escenas[0]
    datos = {f"montaje-1-{k}": foto(f"s1_{k}.jpg") for k in range(len(s1["escena"]["pasos"]) + 1)}
    datos["montaje-3-0"] = foto("s3_0.jpg")
    datos["montaje-1-planta"] = foto("s1_planta.jpg")
    datos["montaje-99-0"] = foto("sesion_inexistente.jpg")                       # se ignora
    datos["montaje-1-50"] = foto("paso_inexistente.jpg")                         # se ignora
    datos["montaje-2-0"] = SimpleUploadedFile("no_es_imagen.jpg", b"<script>", content_type="image/jpeg")  # se ignora
    r = cliente.post(reverse("unidad_docx", args=[ud.pk]), datos)
    assert r.status_code == 200 and r["Content-Type"] == DOCX
    word = Document(io.BytesIO(r.content))
    assert len(word.inline_shapes) == len(s1["escena"]["pasos"]) + 1 + 1 + 1        # sesión 1 (3D, planta, pasos) + sesión 3
    texto = "\n".join(c.text for t in word.tables for fila in t.rows for c in fila.cells)
    texto_u = texto.upper()
    assert "MONTAJE 3D · SESIÓN 1" in texto_u and "MONTAJE 3D · SESIÓN 3" in texto_u and "SESIÓN 2: CÓMO" not in texto_u
    assert f"Paso 1. {s1['escena']['pasos'][0]}" in texto
    assert "Vista de planta" in texto

    # Filtrar solo sesiones 1 y 2 que el docente elija
    datos_filtro = {"sesiones_elegidas": "1,2", "montaje-1-0": foto("s1_0.jpg"), "montaje-2-0": foto("s2_0.jpg")}
    r_filtro = cliente.post(reverse("unidad_docx", args=[ud.pk]), datos_filtro)
    assert r_filtro.status_code == 200 and r_filtro["Content-Type"] == DOCX
    doc_filtro = Document(io.BytesIO(r_filtro.content))
    txt_filtro = "\n".join(c.text for t in doc_filtro.tables for fila in t.rows for c in fila.cells).upper()
    assert "SESIÓN 1 DE 2" in txt_filtro and "SESIÓN 2 DE 2" in txt_filtro
    assert "SESIÓN 3 DE" not in txt_filtro and "SESIÓN 4 DE" not in txt_filtro
    assert "MONTAJE 3D · SESIÓN 1" in txt_filtro and "MONTAJE 3D · SESIÓN 2" in txt_filtro



@pytest.mark.django_db
def test_foto_exige_elegir_habilidad(cliente, catalogo):
    ok, png = cv2.imencode(".png", np.full((120, 160, 3), 60, np.uint8))
    foto = SimpleUploadedFile("foto.png", png.tobytes(), content_type="image/png")
    r = cliente.post(reverse("evaluar"), {"habilidad": "auto", "grado": "7_anos", "motor": "local", "archivo": foto})
    assert r.status_code == 200 and "no es fiable" in r.content.decode()
    assert not Evaluacion.objects.exists()


@pytest.mark.django_db
def test_otro_docente_no_ve_lo_ajeno(cliente, catalogo, django_user_model, tmp_path, django_capture_on_commit_callbacks):
    ev = evaluar(cliente, tmp_path, django_capture_on_commit_callbacks)
    intruso = django_user_model.objects.create_user("intruso", password="x-segura-123")
    cliente.force_login(intruso)
    for url in (reverse("evaluacion", args=[ev.pk]), reverse("fotograma", args=[ev.pk, 1, "imagen"]),
                reverse("reporte", args=[ev.pk]), reverse("evaluacion_estado", args=[ev.pk])):
        assert cliente.get(url).status_code == 404, url
    assert cliente.post(reverse("evaluacion_eliminar", args=[ev.pk])).status_code == 404


@pytest.mark.django_db
def test_eliminar_borra_archivos(cliente, catalogo, tmp_path, django_capture_on_commit_callbacks):
    from pathlib import Path
    ev = evaluar(cliente, tmp_path, django_capture_on_commit_callbacks)
    rutas = [Path(ev.archivo.path)] + [Path(f.imagen.path) for f in ev.fotogramas.all()]
    assert all(p.exists() for p in rutas)
    assert cliente.post(reverse("evaluacion_eliminar", args=[ev.pk])).status_code == 302
    assert not Evaluacion.objects.exists() and not any(p.exists() for p in rutas)


@pytest.mark.django_db
def test_salon_y_evaluacion_grupal(cliente, catalogo, docente, tmp_path, django_capture_on_commit_callbacks):
    r = cliente.post(reverse("grupos"), {"nombre": "2ºB", "grado": "7_anos", "anio": 2026})
    grupo = Grupo.objects.get()
    assert r.status_code == 302 and grupo.institucion is not None          # docente sin institución: se crea una propia

    cliente.post(reverse("grupo", args=[grupo.pk]), {"nombres": "Gómez Pérez, María José\nLuis\n\n"})
    assert set(grupo.estudiantes.values_list("nombres", "apellidos")) == {("María José", "Gómez Pérez"), ("Luis", "")}
    assert cliente.get(reverse("grupo", args=[grupo.pk])).status_code == 200

    est = Estudiante.objects.get(nombres="Luis")
    r = cliente.post(reverse("estudiante_editar", args=[est.pk]), {"nombres": "Luis", "apellidos": "Mora",
                                                                    "consentimiento_video": "on"})
    assert r.status_code == 302 and Estudiante.objects.get(pk=est.pk).consentimiento_video

    r = cliente.post(reverse("grupal_crear", args=[grupo.pk]), {"estudiantes_objetivo": 2})
    grupal = EvaluacionGrupal.objects.get()
    assert r.status_code == 302
    for e in grupo.estudiantes.all():
        evaluar(cliente, tmp_path, django_capture_on_commit_callbacks, estudiante=e.pk, evaluacion_grupal=grupal.pk,
                habilidad="carrera")
    assert grupal.evaluaciones.filter(estado="lista").count() == 2

    html = cliente.get(reverse("grupal", args=[grupal.pk])).content.decode()
    assert "2 de 2 estudiantes" in html and "Matriz de deficiencias" in html

    r = cliente.post(reverse("grupal_plan", args=[grupal.pk]), {"formato": "Retos Cooperativos", "metodologia": "Asignación de Tareas",
                                                                "periodo": 1, "duracion_min": 50, "total_clases": 6})
    ud = UnidadDidactica.objects.get(evaluacion_grupal=grupal)
    assert r.status_code == 302 and ud.habilidad.codigo == "carrera" and ud.sesiones.count() == 6
    assert cliente.get(reverse("estudiante", args=[est.pk])).status_code == 200


@pytest.mark.django_db
def test_asistente_como_la_web(cliente, catalogo, docente, tmp_path, django_capture_on_commit_callbacks):
    """El asistente envía solo habilidad, archivo y observaciones; grado y motor salen de «Configurar clase»."""
    docente.configuracion_clase = {"grado": "9_11_anos", "motor": "gemini"}   # sin clave de Gemini → motor local
    docente.save()
    html = cliente.get(reverse("evaluar")).content.decode()
    assert 'id="wizardStep1"' in html and 'id="uploadZone"' in html and 'class="engine-toggle"' in html
    with django_capture_on_commit_callbacks(execute=True):
        r = cliente.post(reverse("evaluar"), {"habilidad": "salto", "archivo": video_subida(tmp_path)},
                         HTTP_X_REQUESTED_WITH="fetch")
    ev = Evaluacion.objects.get()
    assert r.json() == {"url": reverse("evaluacion", args=[ev.pk]), "estado_url": reverse("evaluacion_estado", args=[ev.pk])}
    assert ev.grado == "9_11_anos" and ev.motor == "local"
    r = cliente.post(reverse("evaluar"), {"habilidad": "salto"}, HTTP_X_REQUESTED_WITH="fetch")
    assert r.status_code == 400 and r.json()["errores"]


@pytest.mark.django_db
def test_barra_del_modo_colectivo(cliente, catalogo, docente):
    docente.configuracion_clase = {"modo": "colectivo", "grado": "8_anos"}
    docente.save()
    assert 'id="groupSetup"' in cliente.get(reverse("evaluar")).content.decode()
    r = cliente.post(reverse("grupal_iniciar"), {"estudiantes_objetivo": 25})
    grupal = EvaluacionGrupal.objects.get()
    assert r["Location"] == f"{reverse('evaluar')}?grupal={grupal.pk}"
    assert grupal.estudiantes_objetivo == 25 and grupal.grupo.grado == "8_anos"
    html = cliente.get(reverse("evaluar")).content.decode()           # el registro sigue abierto sin el parámetro
    assert "Estudiante actual: 1 de 25" in html and f'value="{grupal.pk}"' in html
    cliente.post(reverse("grupal_reiniciar"))
    assert 'id="groupSetup"' in cliente.get(reverse("evaluar")).content.decode()
    cliente.post(reverse("grupal_iniciar"), {"estudiantes_objetivo": 10, "grupo": grupal.grupo.pk})
    assert Grupo.objects.count() == 1 and EvaluacionGrupal.objects.count() == 2

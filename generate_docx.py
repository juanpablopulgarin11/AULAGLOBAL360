import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Configuración de márgenes
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

def add_title(text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    return h

def add_paragraph(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = RGBColor(16, 120, 100)
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.italic = italic
    return p

def add_callout(text, title="¡INFORMACIÓN IMPORTANTE!"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_tit = p.add_run(f"📌 {title}\n")
    r_tit.bold = True
    r_tit.font.size = Pt(10.5)
    r_tit.font.color.rgb = RGBColor(180, 50, 20)
    
    r_body = p.add_run(text)
    r_body.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_bullet(num, title, desc):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_num = p.add_run(f"[{num}] {title}: ")
    r_num.bold = True
    r_num.font.size = Pt(10.5)
    r_num.font.color.rgb = RGBColor(20, 90, 140)
    r_desc = p.add_run(desc)
    r_desc.font.size = Pt(10.5)

def add_image(img_rel_path, caption="", width_inches=6.0):
    full_path = os.path.abspath(img_rel_path)
    if os.path.exists(full_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run()
        r.add_picture(full_path, width=Inches(width_inches))
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(6)
            r_cap = p_cap.add_run(f"Figura: {caption}")
            r_cap.font.size = Pt(9)
            r_cap.italic = True
            r_cap.font.color.rgb = RGBColor(100, 100, 100)
    else:
        p = doc.add_paragraph()
        p.add_run(f"[ Imagen no disponible: {img_rel_path} ]").bold = True

# ==================== PORTADA ====================
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(70)
title_p.paragraph_format.space_after = Pt(12)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = title_p.add_run("AULA GLOBAL 360\n")
r1.bold = True
r1.font.size = Pt(26)
r1.font.color.rgb = RGBColor(10, 80, 140)

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(30)
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = sub_p.add_run("Manual de Usuario y Operación Técnica\nPlataforma de Inteligencia Biomecánica y Curricular")
r2.font.size = Pt(14)
r2.font.color.rgb = RGBColor(70, 70, 70)

info_p = doc.add_paragraph()
info_p.paragraph_format.space_before = Pt(180)
info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_info = info_p.add_run(
    "Enlace Oficial en la Nube (Google Cloud Platform):\n"
    "https://aulaglobal360-528391102571.us-central1.run.app\n\n"
    "Batería de Evaluación de González Palacio & Montoya Grisales\n"
    "Estadios de Desarrollo Motor de David L. Gallahue\n"
    "Educación Física Escolar (5 a 11 Años)"
)
r_info.font.size = Pt(10)
r_info.font.color.rgb = RGBColor(90, 90, 90)

doc.add_page_break()

# ==================== SECCIÓN 1: INTRODUCCIÓN Y ENLACE ====================
add_title("1. Introducción y Enlace de Acceso", level=1)
add_paragraph(
    "Aula Global 360 es una plataforma pedagógica basada en Inteligencia Artificial y Visión por Computadora, "
    "creada específicamente para facilitar y enriquecer la labor de los educadores físicos en las instituciones educativas. "
    "Permite diagnosticar los patrones de movimiento de niños de 5 a 11 años, identificar las fases de desarrollo motor y generar "
    "unidades didácticas personalizadas y listas para el aula de clase.",
    bold_prefix="Propósito Central: "
)
add_callout(
    "Enlace oficial para ingresar a la plataforma:\n"
    "👉 https://aulaglobal360-528391102571.us-central1.run.app\n"
    "No requiere descargas ni configuraciones complejas. Es compatible con computadores, tabletas y teléfonos celulares.",
    title="ENLACE OFICIAL DE ACCESO"
)
add_paragraph("Al abrir el enlace, el sistema te recibirá con la pantalla de inicio institucional:")
add_image("screenshots_auto/cloud_1_landing.png", caption="Pantalla de inicio oficial (Landing Page)")
add_bullet("1", "Menú Institucional", "Enlaces a la fundamentación científica, tecnología de doble motor y habilidades HMB.")
add_bullet("2", "Botón Dashboard", "Acceso directo al entorno de trabajo docente.")

# ==================== SECCIÓN 2: REGISTRO E INICIO DE SESIÓN ====================
doc.add_page_break()
add_title("2. Registro y Gestión de Cuentas Docentes", level=1)
add_paragraph(
    "Para garantizar la privacidad de los estudiantes y permitir que tus planeaciones, salones y evaluaciones se guarden de forma segura, "
    "la plataforma cuenta con un sistema de usuarios docentes."
)
add_title("2.1. Crear una Cuenta de Docente", level=2)
add_paragraph(
    "El registro es completamente libre y toma menos de un minuto. Solo debes ingresar a la sección de registro "
    "(https://aulaglobal360-528391102571.us-central1.run.app/cuentas/registro/):"
)
add_image("screenshots_auto/cloud_2_registro.png", caption="Formulario de creación de cuenta de docente")
add_bullet("1", "Nombre de usuario", "Un alias que recordarás fácilmente (ej: profe_carlos).")
add_bullet("2", "Nombres y Apellidos", "Tus datos personales con los cuales se firmarán los reportes y planeaciones en Word.")
add_bullet("3", "Correo Electrónico", "Tu dirección de contacto institucional o personal.")
add_bullet("4", "Contraseña", "Clave de seguridad de mínimo 8 caracteres para proteger tu cuenta.")

add_title("2.2. Iniciar Sesión", level=2)
add_paragraph(
    "Cuando ya dispongas de tu cuenta, ingresa desde https://aulaglobal360-528391102571.us-central1.run.app/cuentas/ingresar/ "
    "con tu usuario y contraseña:"
)
add_image("screenshots_auto/cloud_3_ingresar.png", caption="Pantalla de inicio de sesión de usuario")

# ==================== SECCIÓN 3: LA NUEVA BARRA DE NAVEGACIÓN ====================
doc.add_page_break()
add_title("3. La Barra de Navegación del Docente", level=1)
add_paragraph(
    "Una vez que inicias sesión, en la parte superior de la aplicación encontrarás la barra principal con los 4 módulos operativos:"
)
add_image("screenshots_auto/cloud_header_full_hd.png", caption="Barra superior completa en alta definición (Navegación, Motor y Configuración)", width_inches=6.2)
add_image("screenshots_auto/cloud_header_left_hd.png", caption="Detalle nítido de los 4 módulos: Evaluar, Historial, Salones y Planeaciones", width_inches=4.8)
add_bullet("1", "Evaluar", "Módulo principal para cargar videos/fotos de los estudiantes y obtener el diagnóstico biomecánico inmediato.")
add_bullet("2", "Historial", "Consulta el registro acumulativo de todas las evaluaciones previas, resultados y diagnósticos guardados.")
add_bullet("3", "Salones", "Administra los grados escolares (ej: 1A, 2B), las listas de alumnos y el estado de los consentimientos informados (Ley 1581).")
add_bullet("4", "Planeaciones", "Accede a las unidades didácticas generadas, visualiza las sesiones curriculares con montaje 3D y descárgalas en Word.")

add_paragraph(
    "Adicionalmente, en el centro de la barra superior encontrarás el selector de motor de análisis (Análisis rápido local vs Análisis con IA) "
    "y en la esquina derecha el botón verde 'Configurar clase'."
)
add_image("screenshots_auto/cloud_nav_evaluar.png", caption="Pantalla de trabajo general con barra de navegación y panel de evaluación")

# ==================== SECCIÓN 4: CONFIGURACIÓN DE CLASE ====================
doc.add_page_break()
add_title("4. Panel de Configuración de Clase (El Cerebro de la IA)", level=1)
add_paragraph(
    "Antes de subir un video, es indispensable ajustar el contexto escolar presionando el botón 'Configurar clase' en la esquina superior derecha. "
    "Esto despliega un panel lateral donde personalizas todos los parámetros de la sesión:"
)
add_image("screenshots_auto/cloud_nav_config_drawer.png", caption="Panel lateral Drawer para configuración de clase")
add_bullet("1", "Edad y Grado", "Transición a 5º grado (5 a 11 años). La IA compara el movimiento contra los hitos maduros correspondientes a la edad.")
add_bullet("2", "Formato Didáctico", "Cuento Motor (narrativo para pequeños), Circuito de Estaciones (postas por grupos), Retos Cooperativos o Juego Libre.")
add_bullet("3", "Metodología", "Descubrimiento Guiado (preguntas inductivas), Resolución de Problemas o Asignación de Tareas (mando directo).")
add_bullet("4", "Tiempos y Frecuencia", "Duración de la clase (minutos) y cantidad de clases del período. La IA calcula el tiempo exacto de cada fase.")
add_bullet("5", "Materiales Disponibles", "Filtro estricto: desmarca los materiales que no tengas. La IA NUNCA sugerirá un material que no esté disponible.")
add_bullet("6", "Modo de Evaluación", "Individual (un estudiante) o Colectivo (salón completo con plan consolidado).")
add_bullet("7", "Motor de Análisis", "Elige entre 'Análisis rápido' (ejecutado en el servidor) o 'Análisis detallado' (visión avanzada con Google Gemini).")

# ==================== SECCIÓN 5: FLUJO DE EVALUACIÓN PASO A PASO ====================
doc.add_page_break()
add_title("5. Flujo de Trabajo y Evaluación Paso a Paso", level=1)
add_paragraph("El proceso de diagnóstico se desarrolla en sencillos pasos visuales:")

add_title("Paso 1: Elección de Habilidad Motriz Básica (HMB)", level=2)
add_image("screenshots_auto/auto_4_paso1.png", caption="Paso 1: Detección automática o selección manual de la habilidad")
add_bullet("1", "Detección Automática", "El sistema analiza el video e infiere automáticamente la habilidad que el niño está ejecutando.")
add_bullet("2", "Selección Manual", "El docente puede forzar la habilidad: Carrera, Salto horizontal, Lanzamiento, Atrapar, Patear, Equilibrio, etc.")

add_title("Paso 2: Consulta de la Guía de Grabación", level=2)
add_paragraph("Antes de subir el archivo, presiona el botón 'Ver guía de grabación' para verificar las recomendaciones biomecánicas:")
add_image("screenshots_auto/auto_4_paso2_btn_guia.png", caption="Botón para desplegar la guía de grabación recomendada")
add_image("screenshots_auto/auto_4_paso3_modal.png", caption="Cuadro modal con indicaciones de distancia, encuadre y planos corporales")
add_bullet("1", "Duración", "3 a 5 segundos de grabación son suficientes. Grabar 1 segundo antes de la acción y 1 segundo después.")
add_bullet("2", "Encuadre", "Cuerpo completo de pies a cabeza a 3 o 4 metros de distancia.")
add_bullet("3", "Plano Anatómico", "Perfil (lateral) para carreras y saltos; Frente (frontal) para lanzamientos, recepciones y equilibrio.")

add_title("Paso 3: Subida del Video o Fotografía", level=2)
add_paragraph("Arrastra el video al recuadro o haz clic en él para seleccionarlo desde la memoria de tu dispositivo:")
add_image("screenshots_auto/auto_4_paso4_upload.png", caption="Zona interactiva de subida de archivos")

add_title("Paso 4: Extracción de Fotogramas Clave", level=2)
add_paragraph("Al cargar el video, el sistema extraerá los fotogramas clave y trazará el esqueleto de articulaciones:")
add_image("screenshots_auto/auto_4_paso5_fotogramas.png", caption="Tira de fotogramas clave con esqueleto biomecánico")
add_callout(
    "Espera unos segundos a que la tira de imágenes aparezca en pantalla. "
    "El botón verde 'Analizar movimiento' se activará cuando los fotogramas estén listos. "
    "Solo entonces haz clic para iniciar el diagnóstico.",
    title="REGLA DE ORO DE PROCESAMIENTO"
)

add_title("Paso 5: Resultados del Diagnóstico", level=2)
add_image("screenshots_auto/auto_5_resultados.png", caption="Contenedor de resultados con puntaje y estadio de Gallahue")
add_bullet("1", "Porcentaje de Madurez", "Puntaje global sobre 100% alcanzado por el estudiante.")
add_bullet("2", "Estadio de Gallahue", "Inicial (patrón inmaduro), Elemental (en desarrollo) o Maduro (óptimo biomecánicamente).")
add_bullet("3", "Criterios Biomecánicos", "Detalle de los ángulos, balanceo de brazos y apoyos superados y por corregir.")

# ==================== SECCIÓN 6: SALONES, HISTORIAL Y PLANEACIONES ====================
doc.add_page_break()
add_title("6. Salones, Historial y Banco de Planeaciones", level=1)

add_title("6.1. Salones de Clase y Evaluación Grupal", level=2)
add_paragraph(
    "En la pestaña 'Salones' puedes crear tus cursos escolares y matricular a tus estudiantes. "
    "El sistema gestiona los consentimientos informados de los acudientes según la Ley 1581 de protección de datos de menores. "
    "Al realizar evaluaciones de un grupo completo, la plataforma consolida las debilidades comunes y elabora un plan grupal unificado."
)
add_image("screenshots_auto/cloud_nav_salones.png", caption="Módulo de gestión de salones y grupos escolares")

add_title("6.2. Historial de Evaluaciones", level=2)
add_paragraph(
    "En la pestaña 'Historial' podrás consultar en cualquier momento los diagnósticos efectuados, filtrar por fecha o habilidad "
    "y verificar el progreso motriz de cada estudiante a través del tiempo."
)
add_image("screenshots_auto/cloud_nav_historial.png", caption="Módulo de historial de evaluaciones realizadas")

add_title("6.3. Banco de Planeaciones Curriculares", level=2)
add_paragraph(
    "En la pestaña 'Planeaciones' se almacenan todas las unidades didácticas generadas para tus clases. "
    "Cada planeación cuenta con sesiones editables y la opción de montaje espacial tridimensional."
)
add_image("screenshots_auto/cloud_nav_planeaciones.png", caption="Módulo de planeaciones y unidades didácticas")

# ==================== SECCIÓN 7: EXPORTACIÓN A WORD ====================
doc.add_page_break()
add_title("7. Exportación a Microsoft Word (.docx)", level=1)
add_paragraph(
    "Una de las mayores ventajas de Aula Global 360 es la generación de documentos oficiales listos para imprimir o anexar a tu plan de área:"
)
add_bullet("1", "Reporte Biomecánico Individual (.docx)", "Documento formal con los datos del estudiante, puntaje, estadio de Gallahue, tabla de criterios y espacio para firma docente. Ideal para entregar a padres de familia o para el historial pedagógico.")
add_bullet("2", "Planeación de Clase Completa (.docx)", "Unidad didáctica estructurada con objetivos, distribución del tiempo (inicio, desarrollo, final), consignas verbales ('El Lenguaje del Profe'), medidas de seguridad y adaptaciones inclusivas DUA.")

# ==================== SECCIÓN 8: RECOMENDACIONES FINALES ====================
doc.add_page_break()
add_title("8. Recomendaciones Finales para el Docente", level=1)
add_bullet("1", "Uso directo en la cancha", "La plataforma es 100% responsiva. Puedes abrirla en tu teléfono móvil durante la clase, grabar al estudiante y mostrarle inmediatamente el esqueleto de su movimiento para retroalimentarlo.")
add_bullet("2", "Iluminación adecuada", "Realiza las tomas en espacios bien iluminados y con ropa deportiva de contraste para que la inteligencia artificial detecte con máxima precisión las articulaciones corporales.")
add_bullet("3", "Privacidad y Seguridad", "Los videos procesados son confidenciales y solo pueden ser consultados por el docente titular de la cuenta.")

output_path_v11 = os.path.abspath("Manual_Usuario_AulaGlobal360_v11.docx")
output_path_main = os.path.abspath("Manual_Usuario_AulaGlobal360.docx")

doc.save(output_path_v11)
doc.save(output_path_main)
print(f"Documentos guardados con éxito:\n- {output_path_v11}\n- {output_path_main}")
